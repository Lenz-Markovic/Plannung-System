"""Excel "Bestellliste" of "💶 Material & Kosten" (openpyxl)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .rules.material import DUE, GROUP_LABELS, OPEN, PLANNED

EUR = '#,##0.00 "€"'
HEAD = PatternFill("solid", fgColor="1F4E8C")
TOTAL = PatternFill("solid", fgColor="E8EEF7")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))


def material_workbook(summary, labels):
    """labels: category code -> name. Returns the bytes of the .xlsx file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bestellliste"
    ws["A1"] = f"Material & Kosten Montage {summary.start:%d.%m.%Y} – {summary.end:%d.%m.%Y}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "gezählt: " + " + ".join(GROUP_LABELS[g] for g in summary.counted)
    ws["A2"].font = Font(italic=True, color="555555")

    head = ["Art.-Nr.", "Bezeichnung", "Kategorie", "geplant", "fällig ungeplant", "offen ohne Frist",
            "Stück gesamt", "Preis/Stück", "Summe", "zuerst gebraucht", "Aufträge"]
    ws.append([])
    ws.append(head)
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in summary.rows:
        if not row.count(summary.counted):
            continue
        ws.append([row.article, row.description, labels.get(row.category, row.category or "–"),
                   row.pieces[PLANNED], row.pieces[DUE], row.pieces[OPEN], row.count(summary.counted),
                   float(row.price), float(row.money(summary.counted)),
                   row.first_needed, ", ".join(sorted(row.orders))])
        line = ws.max_row
        ws.cell(line, 8).number_format = ws.cell(line, 9).number_format = EUR
        ws.cell(line, 10).number_format = "DD.MM.YYYY"
        for cell in ws[line]:
            cell.border = THIN
    ws.append(["", "Gesamt", "", summary.pieces[PLANNED], summary.pieces[DUE], summary.pieces[OPEN],
               summary.total_pieces, "", float(summary.total), "", f"{summary.total_orders} Aufträge"])
    last = ws.max_row
    for cell in ws[last]:
        cell.font, cell.fill = Font(bold=True), TOTAL
    ws.cell(last, 9).number_format = EUR
    for column, width in zip("ABCDEFGHIJK", [11, 38, 24, 9, 10, 10, 9, 11, 13, 12, 40]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"

    weeks = wb.create_sheet("Budget je Woche")
    weeks.append(["Woche ab", "Betrag"])
    for cell in weeks[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
    for monday, money in summary.by_week:
        weeks.append([monday, float(money)])
        weeks.cell(weeks.max_row, 1).number_format = "DD.MM.YYYY"
        weeks.cell(weeks.max_row, 2).number_format = EUR
    weeks.column_dimensions["A"].width, weeks.column_dimensions["B"].width = 14, 14

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
