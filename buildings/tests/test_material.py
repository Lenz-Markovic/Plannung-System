"""💶 Material & Kosten page: only Admin, window, prices, Excel."""

import datetime
import io
from decimal import Decimal

import pytest
from django.contrib.auth.models import Group, User
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from buildings.material import material_summary, parse_price
from buildings.models import ArticlePrice, DeviceCategory, InstallationOrder
from core import roles
from planning.models import StopKind, TourStop

pytestmark = pytest.mark.django_db
URL = "orders:material"


@pytest.fixture
def demo(demo_import, monkeypatch):
    monkeypatch.setattr(timezone, "localdate", lambda *args: datetime.date(2026, 9, 28))


def login(role, name):
    user = User.objects.create_user(username=name)
    user.groups.add(Group.objects.get(name=role))
    client = Client()
    client.force_login(user)
    return client


def a_planned_installation():
    return TourStop.objects.filter(kind=StopKind.INSTALLATION).select_related("tour", "installation_order").order_by("tour__date").first()


def test_only_admin_sees_money(demo):
    assert login(roles.ADMIN, "admin").get(reverse(URL)).status_code == 200
    for role in (roles.DISPATCHER, roles.PROCESSING, roles.READER, roles.MANAGEMENT):
        assert login(role, f"u-{role}").get(reverse(URL)).status_code == 403
    dispo = login(roles.DISPATCHER, "dispo2")
    assert "Material &amp; Kosten" not in dispo.get(reverse("orders:list")).content.decode()


def test_planned_installation_is_in_its_window(demo):
    stop = a_planned_installation()
    day = stop.tour.date
    s = material_summary(day, day)
    assert stop.installation_order.re_number in {u.order for r in s.rows for u in r.uses}
    expected = sum(i.quantity * (i.category.price if i.category else 0) for i in stop.installation_order.items.all())
    assert s.money["planned"] >= expected > 0
    page = login(roles.ADMIN, "admin").get(reverse(URL), {"z": "frei", "von": day.isoformat(), "bis": day.isoformat()}).content.decode()
    assert "Bestellliste" in page and stop.installation_order.items.first().article_number in page


def test_presets_and_open_switch(demo):
    admin = login(roles.ADMIN, "admin")
    page = admin.get(reverse(URL), {"z": "1m"}).content.decode()
    assert "28.09. – 27.10.2026" in page
    s_without = material_summary(datetime.date(2026, 9, 28), datetime.date(2026, 10, 27))
    s_with = material_summary(datetime.date(2026, 9, 28), datetime.date(2026, 10, 27), with_open=True)
    assert s_with.total >= s_without.total


def test_prices_can_be_changed(demo):
    admin = login(roles.ADMIN, "admin")
    admin.post(reverse("orders:material_price"), {"category": "EHKV", "price": "13,50", "z": "2w"})
    assert DeviceCategory.objects.get(code="EHKV").price == Decimal("13.50")
    response = admin.post(reverse("orders:material_price"), {"article": "322011F", "description": "EHKV", "price": "11,20", "z": "2w"})
    assert ArticlePrice.objects.get(article_number="322011F").price == Decimal("11.20")
    assert "Artikelpreis 322011F gespeichert" in response.content.decode()
    admin.post(reverse("orders:material_price"), {"article": "322011F", "price": "", "z": "2w"})  # back to category
    assert not ArticlePrice.objects.exists()
    bad = admin.post(reverse("orders:material_price"), {"category": "EHKV", "price": "zwölf"})
    assert bad["HX-Reswap"] == "none" and "kein Preis" in bad.content.decode()


def test_parse_price():
    assert parse_price("1.250,50 €") == Decimal("1250.50") and parse_price("12.5") == Decimal("12.50")
    assert parse_price("") is None
    with pytest.raises(ValueError):
        parse_price("-3")


def test_excel_order_list(demo):
    stop = a_planned_installation()
    day = stop.tour.date.isoformat()
    response = login(roles.ADMIN, "admin").get(reverse(URL), {"z": "frei", "von": day, "bis": day, "format": "xlsx"})
    assert response["Content-Disposition"].startswith('attachment; filename="Material_')
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb["Bestellliste"]
    assert ws["A4"].value == "Art.-Nr." and ws.cell(ws.max_row, 2).value == "Gesamt"
    assert "Budget je Woche" in wb.sheetnames


def test_done_orders_without_plan_are_not_bought_again(demo):
    order = InstallationOrder.objects.filter(tour_stops__isnull=True).first()
    order.status = "done"
    order.save()
    s = material_summary(datetime.date(2026, 1, 1), datetime.date(2027, 12, 31), with_open=True)
    assert order.re_number not in {u.order for r in s.rows for u in r.uses}


def test_per_contract_view_and_excel_sheet(demo):
    stop = a_planned_installation()
    day = stop.tour.date.isoformat()
    admin = login(roles.ADMIN, "admin")
    params = {"z": "frei", "von": day, "bis": day, "ansicht": "auftrag"}
    page = admin.get(reverse(URL), params).content.decode()
    order = stop.installation_order
    assert "Je Auftrag ·" in page and order.re_number in page and order.street in page
    response = admin.get(reverse(URL), {**params, "format": "xlsx"})
    ws = load_workbook(io.BytesIO(response.content))["Je Auftrag"]
    assert ws["C1"].value == "RE-Nr." and order.re_number in [c.value for c in ws["C"]]
    # a price change keeps the chosen view
    kept = admin.post(reverse("orders:material_price"), {"category": "EHKV", "price": "12,00", **params}).content.decode()
    assert "Je Auftrag ·" in kept
