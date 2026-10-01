"""
Excel export of tours, in exactly the layout of the prototype's
"Fahrpläne … je Monteur" workbook (excelFahrplaene / blattListe / blattTag):

- sheet "Fahrpläne": all tours one below the other          (only if > 1 tour)
- one sheet per person: his/her tours one below the other   (only if > 1 tour)
- one print sheet per tour: A4 portrait, big font, a block per building

Column widths, fonts, colours, borders, row heights and page setup are taken
over 1:1 from the prototype (ExcelJS there, openpyxl here).
"""

import datetime
import math
import re
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from buildings.display import INSTALLATION_SHORT
from buildings.rules.file_numbers import extract_re_numbers
from buildings.rules.reading_time import is_manual_reading

from .models import RoutingSource, StopKind, TourStatus
from .rules.working_time import MAX_NET_MINUTES, schedule_day

# --- styles (XF / XA / XK / XC / RAHMEN in the prototype) ---------------------
NAVY, GREY, WARN = "FF1F3864", "FFD9D9D9", "FFFFE699"
THIN = Side(style="thin", color="FF000000")
MEDIUM = Side(style="medium", color="FF000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WEEKDAYS = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]
MONTHS = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]
MONTHS_SHORT = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]
HEADINGS = ["Uhrzeit ab:", "AZ", "Adresse", "RE", "ToDo", "FUNK", "Zur Anmeldung", "Termin bestätigt"]
DEVICE_NAMES = {"HKV": "HKV", "WMZ": "Wärmezähler", "WWZ": "Warmwasserzähler", "KWZ": "Kaltwasserzähler", "RWM": "Rauchwarnmelder"}


def fill(argb):
    return PatternFill("solid", fgColor=argb)


def arial(**kwargs):
    return Font(name="Arial", size=kwargs.pop("size", 10), **kwargs)


def calibri(**kwargs):
    return Font(name="Calibri", size=kwargs.pop("size", 14), **kwargs)


def grey_text():
    return calibri(size=9, italic=True, color="FF595959")


def long_date(date):
    """'Montag 07.12.26' (datumLang in the prototype)."""
    return f"{WEEKDAYS[date.weekday()]} {date:%d.%m.%y}"


def hh_mm(minutes):
    return f"{int(minutes) // 60:02d}:{round(minutes) % 60:02d}"


def minutes_of(clock):
    return clock.hour * 60 + clock.minute


# --- data of one tour, prepared for the sheets --------------------------------

@dataclass
class StopRow:
    start: str
    end: str
    address: str
    zip_code: object
    city: str
    az: object
    re: object
    stichtag: str
    todo: list
    anlage: str
    hint: str
    highlight: bool           # yellow ToDo cell: apartment access needed
    drive_after: int | None


@dataclass
class TourSheet:
    date: datetime.date
    person: str
    stops: list = field(default_factory=list)
    work: int = 0
    installation_work: int = 0  # part of `work` that is installation (Montage)
    drive: int = 0
    break_after: int | None = None
    break_minutes: int = 0
    provisional: bool = False
    source: str = ""
    commute: int | None = None

    @property
    def net(self):
        return self.work + self.drive

    def net_text(self):
        """nettoText() in the prototype."""
        reading = self.work - self.installation_work
        # Same words as the prototype for reading tours; Montage named where it is part of the day
        parts = ([f"Ablesung {reading} min"] if reading or not self.installation_work else []) + (
            [f"Montage {self.installation_work} min"] if self.installation_work else [])
        text = f"Netto-Arbeitszeit: {hh_mm(self.net)}   ({' + '.join(parts)} + Fahrzeit im Tag {self.drive} min)"
        if self.net > MAX_NET_MINUTES:
            text += f"   |  {self.net - MAX_NET_MINUTES} min über 7,5 h"
        if self.provisional:
            text += "   |  VORLÄUFIG – noch ohne TomTom-Zeiten"
        return text

    def towns(self):
        """'Nagold → Calw → Böblingen' (orteKette)."""
        towns = []
        for stop in self.stops:
            if stop.city and (not towns or towns[-1] != stop.city):
                towns.append(stop.city)
        return " → ".join(towns)


def _devices(building):
    """'50× HKV', '1× Wärmezähler', ... (geraete + vorlageDaten in the prototype)."""
    items = [(building.hkv_count, "HKV"), (building.wmz_count, "WMZ"), (building.wwz_count, "WWZ"), (building.kwz_count, "KWZ")]
    result = [f"{count}× {DEVICE_NAMES[code]}" for count, code in items if count]
    if building.has_rwm:
        result.append((f"{building.rwm_count}× " if building.rwm_count else "") + DEVICE_NAMES["RWM"])
    return result


def _stichtag(building):
    text = " ".join([building.remark, building.note, building.handwritten_note])
    match = re.search(r"Stichtag\s*:?\s*(\d{1,2}\.\d{1,2}\.)", text, re.IGNORECASE)
    return match.group(1) if match else f"{building.stichtag:%d.%m.}"


def _re_value(numbers):
    numbers = sorted(n.upper().removeprefix("RE") for n in numbers)
    if not numbers:
        return "—"
    if len(numbers) == 1:
        return int(numbers[0]) if numbers[0].isdigit() else numbers[0]
    return numbers[0] + " / RE" + " / RE".join(numbers[1:])


def stop_row(stop, start, end):
    """What the sheets show for one stop (vorlageDaten in the prototype)."""
    building, order = stop.building, stop.installation_order
    target = building or order
    access = building.access if building and stop.kind != StopKind.NOTICE else None  # hanging a notice: no flat
    core = building.file_number_core if building else order.building_file_number_core
    az = int(core) if core.isdigit() else (building.file_number if building else order.building_file_number)
    if stop.kind == StopKind.READING:
        head = "Ablesung" + (" manuell" if is_manual_reading(building.reading_type)
                             else (" Funk" if re.search(r"Funk|^F", building.reading_type or "", re.I) else ""))
        todo = [head]
        re_numbers = extract_re_numbers(building.order_reference, building.handwritten_note)
    elif stop.kind == StopKind.NOTICE:
        from documents.notices import appointment_of
        from documents.notice_rules import units_text

        appointment = appointment_of(stop)
        todo = ["📄 Aushang aufhängen" + (f" für {'Montage' if order else 'Ablesung'} am {appointment.tour.date:%d.%m.%Y}"
                                          if appointment else "")]
        if appointment and units_text(appointment.notice_scope, appointment.notice_units):
            todo.append(units_text(appointment.notice_scope, appointment.notice_units))
        re_numbers = {order.re_number} if order else set()
    elif stop.kind == StopKind.HELP:
        helped = stop.help_tour.people_label if stop.help_tour else "?"
        todo = [f"🤝 Hilfe bei {helped}: " + (f"Montage {order.re_number}" if order else "Ablesung")]
        re_numbers = {order.re_number} if order else extract_re_numbers(building.order_reference, building.handwritten_note)
    else:
        todo = [f"Montage {order.re_number}"] + [part.strip() for part in order.summary.split(",") if part.strip()]
        re_numbers = {order.re_number}
    if access and access.important:
        extra = [r for r in access.reasons if not r.startswith("Ableseart")]
        todo.append("🔑 Zugang Wohnung" + (" " + ", ".join(access.units) if access.units else "")
                    + (": " + ", ".join(extra) if extra else ""))
    if access and access.room:
        todo.append("🚪 Heizraum/Keller")
    if access and "RWM-Prüfung in den Wohnungen" in access.reasons:
        todo.append("RWM-Prüfung")
    if stop.kind == StopKind.READING:
        todo += _devices(building)
    last = getattr(stop, "last_visit", None)
    if getattr(stop, "attempt", 1) > 1:  # 🔁 Nachtermin: which visit, and what is still to do from last time
        todo.insert(0, f"🔁 {stop.attempt}. Termin" + (f" – noch zu tun: {last.todo}" if last and last.todo else ""))

    hints = []
    if building and building.remark:
        hints.append(building.remark)
    if access and access.key_hint and access.key_hint not in (building.remark or ""):
        hints.append(f"Schlüssel: {access.key_hint}")
    hints += [f"⚠ {warning}" for warning in (stop.route_warnings or [])]

    zip_code = target.zip_code
    return StopRow(
        start=start, end=end, address=target.street,
        zip_code=int(zip_code) if str(zip_code).isdigit() else zip_code, city=target.city,
        az=az, re=_re_value(re_numbers),
        stichtag=_stichtag(building) if building else (order.raw_data or {}).get("stichtag") or "—",
        todo=todo, anlage=INSTALLATION_SHORT.get(building.installation_type, "") if building else "",
        hint=" · ".join(hints), highlight=bool(access and access.important),
        drive_after=stop.drive_to_next_minutes,
    )


def tour_sheet(tour):
    stops = list(tour.stops.select_related("building", "installation_order", "help_tour__employee").order_by("position"))
    from .visits import attach_attempts
    attach_attempts(stops, tour.date)
    work = [s.work_minutes for s in stops]
    drives = [s.drive_to_next_minutes or 0 for s in stops[:-1]]
    # Old/imported tours may have no times yet: calculate them like the plan does.
    plan = schedule_day(tour.start_time, work, drives, tour.break_minutes or 30)
    sheet = TourSheet(
        date=tour.date, person=tour.people_label, work=sum(work), drive=sum(drives),
        installation_work=sum(s.work_minutes for s in stops if s.kind == StopKind.INSTALLATION),
        provisional=tour.status == TourStatus.PROVISIONAL or tour.needs_recalculation or tour.routing_source != RoutingSource.TOMTOM,
        source=tour.routing_note, commute=tour.commute_to_minutes,
    )
    if tour.break_after_position:
        sheet.break_after, sheet.break_minutes = tour.break_after_position - 1, tour.break_minutes
    elif plan.break_after_index is not None:
        sheet.break_after, sheet.break_minutes = plan.break_after_index, plan.break_minutes
    for stop, times in zip(stops, plan.stops):
        start, end = stop.start_time or times.start, stop.end_time or times.end
        sheet.stops.append(stop_row(stop, f"{start:%H:%M}", f"{end:%H:%M}"))
    if sheet.stops:
        sheet.stops[-1].drive_after = None  # no drive after the last stop
    return sheet


# --- cell helpers ---------------------------------------------------------------

def _field(ws, r1, c1, r2, c2, value, font=None, align=None, border=None):
    """Write a (merged) field like feld() in the prototype."""
    cell = ws.cell(r1, c1)
    cell.value = value
    if font:
        cell.font = font
    cell.alignment = Alignment(**{"vertical": "center", "wrap_text": True, **(align or {})})
    if border:
        for row in range(r1, r2 + 1):
            for column in range(c1, c2 + 1):
                ws.cell(row, column).border = border
    if r2 > r1 or c2 > c1:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    return cell


def _height(ws, row, height):
    if height:
        ws.row_dimensions[row].height = height


def _page(ws, landscape, last_row, margins):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_area = f"A1:H{max(1, last_row)}"
    ws.page_margins = PageMargins(**margins)


# --- print sheet per tour (blattTag) --------------------------------------------

def day_sheet(wb, sheet, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, width in enumerate([11, 10.29, 25.14, 12.57, 21.29, 8.71, 14.14, 14.14]):
        ws.column_dimensions["ABCDEFGH"[i]].width = width
    bottom = Border(bottom=MEDIUM)
    z = 1
    _height(ws, z, 26.1)
    _field(ws, z, 1, z, 1, "Datum:", calibri(size=16, bold=True), border=bottom)
    _field(ws, z, 2, z, 4, long_date(sheet.date), calibri(size=16, bold=True), border=bottom)
    _field(ws, z, 5, z, 6, "Mitarbeiter:", calibri(size=16, bold=True), {"horizontal": "right"}, bottom)
    _field(ws, z, 7, z, 8, sheet.person, calibri(size=16, bold=True), {"horizontal": "right"}, bottom)
    z += 1
    _height(ws, z, 20.1)
    _field(ws, z, 1, z, 8, sheet.towns(), calibri(italic=True, color="FF595959"), border=Border(top=THIN))
    z += 1
    _height(ws, z, 6)
    z += 1
    if sheet.commute:
        _height(ws, z, 15)
        _field(ws, z, 1, z, 8, f"Anfahrt von zu Hause: {sheet.commute} Min.", grey_text())
        z += 1

    for i, s in enumerate(sheet.stops):
        _height(ws, z, 39.95)
        for j, text in enumerate(HEADINGS, start=1):
            _field(ws, z, j, z, j, text, calibri(bold=True), {"horizontal": "center"}, BOX).fill = fill(GREY)
        z += 1
        todo_lines = sum(max(1, math.ceil(len(t) / 16)) for t in s.todo)
        height = max(21.95, math.ceil((todo_lines * 21 + 8) / 3))
        for k in range(3):
            _height(ws, z + k, max(height, 39.95 if len(str(s.city)) > 18 else 0) if k == 2 else height)
        _field(ws, z, 1, z + 2, 1, f"{s.start}\n-\n{s.end}", calibri(bold=True), {"horizontal": "center"}, BOX)
        _field(ws, z, 2, z, 2, s.az, calibri(bold=True), {"horizontal": "center"}, BOX)
        _field(ws, z + 1, 2, z + 1, 2, "Handy", calibri(bold=True), {"horizontal": "center"}, BOX)
        _field(ws, z + 2, 2, z + 2, 2, "☐ ja", calibri(), {"horizontal": "center"}, BOX)
        _field(ws, z, 3, z, 3, s.address, calibri(bold=True), {"horizontal": "left"}, BOX)
        _field(ws, z + 1, 3, z + 1, 3, s.zip_code, calibri(), {"horizontal": "left"}, BOX)
        _field(ws, z + 2, 3, z + 2, 3, s.city, calibri(), {"horizontal": "left"}, BOX)
        _field(ws, z, 4, z, 4, s.re, calibri(), {"horizontal": "center"}, BOX)
        _field(ws, z + 1, 4, z + 1, 4, "Stichtag", calibri(bold=True), {"horizontal": "center"}, BOX)
        _field(ws, z + 2, 4, z + 2, 4, s.stichtag, calibri(), {"horizontal": "center"}, BOX)
        todo = _field(ws, z, 5, z + 2, 5, "\n".join(s.todo), calibri(), {"vertical": "top"}, BOX)
        if s.highlight:
            todo.fill = fill(WARN)
        _field(ws, z, 6, z + 2, 6, s.anlage, calibri(size=11), {"horizontal": "center"}, BOX)
        _field(ws, z, 7, z + 2, 7, "", calibri(), border=BOX)
        _field(ws, z, 8, z + 2, 8, "", calibri(), border=BOX)
        z += 3
        if s.hint:
            _height(ws, z, max(1, math.ceil(len(s.hint) / 110)) * 12 + 4)
            _field(ws, z, 1, z, 8, s.hint, grey_text(), {"vertical": "top"})
            z += 1
        _height(ws, z, 21.95)
        ws.merge_cells(start_row=z, start_column=1, end_row=z, end_column=8)
        z += 1
        if sheet.break_after == i and sheet.break_minutes:
            end = minutes_of(datetime.time.fromisoformat(s.end))
            _height(ws, z, 15)
            _field(ws, z, 1, z, 8, f"Mittagspause {hh_mm(end)}–{hh_mm(end + sheet.break_minutes)}", grey_text())
            z += 1
        if i < len(sheet.stops) - 1 and s.drive_after is not None:
            _height(ws, z, 15)
            _field(ws, z, 1, z, 8, f"Fahrtzeit: {s.drive_after} Min.", grey_text())
            z += 1
    _height(ws, z, 6)
    z += 1
    _height(ws, z, 20.1)
    red = sheet.net > MAX_NET_MINUTES or sheet.provisional
    _field(ws, z, 1, z, 8, sheet.net_text(), calibri(italic=True, color="FFC00000" if red else "FF404040"), border=Border(top=THIN))
    z += 1
    _height(ws, z, 12.95)
    _field(ws, z, 1, z, 8, f"Fahrzeiten: {sheet.source}", calibri(size=8, italic=True, color="FF808080"))
    _page(ws, landscape=False, last_row=z, margins=dict(left=0.47, right=0.39, top=0.6, bottom=0.47, header=0.25, footer=0.25))
    ws.oddHeader.left.text, ws.oddHeader.left.font, ws.oddHeader.left.size, ws.oddHeader.left.color = "Seite &P von &N", "Calibri", 12, "808080"
    return ws


# --- list sheet: tours one below the other (blattListe) ------------------------

def list_sheet(wb, name, sheets, title=None):
    ws = wb.create_sheet(name)
    if title:
        ws.freeze_panes = "A2"
    for i, width in enumerate([13, 10, 30, 13, 26, 10, 15, 16]):
        ws.column_dimensions["ABCDEFGH"[i]].width = width
    z = 1

    def navy_row(text):
        nonlocal z
        for j in range(1, 9):
            cell = ws.cell(z, j)
            cell.fill = fill(NAVY)
            cell.font = arial(size=12, bold=True, color="FFFFFFFF")
            cell.border = Border(bottom=MEDIUM)
        ws.cell(z, 1).value = text
        ws.row_dimensions[z].height = 21.95
        z += 1

    def box_row(values, fonts=None):
        nonlocal z
        for j in range(1, 9):
            cell = ws.cell(z, j)
            cell.border = BOX
            cell.font = (fonts or {}).get(j) or arial()
            cell.alignment = Alignment(vertical="center", wrap_text=j in (3, 5))
            if values.get(j) is not None:
                cell.value = values[j]
        z += 1

    if title:
        navy_row(title)
        z += 1
    for sheet in sheets:
        navy_row(f"Datum: {long_date(sheet.date)}      Mitarbeiter: {sheet.person}")
        ws.cell(z, 1).value = (f"Anfahrt von zu Hause: {sheet.commute} min (vor Terminbeginn, keine Arbeitszeit)"
                               if sheet.commute else sheet.towns())
        ws.cell(z, 1).font = arial(size=8, italic=True)
        z += 1
        for j, text in enumerate(HEADINGS, start=1):
            cell = ws.cell(z, j)
            cell.value, cell.font, cell.fill, cell.border = text, arial(bold=True), fill(GREY), BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        z += 1
        for i, s in enumerate(sheet.stops):
            bold = arial(bold=True)
            for k in range(max(3, len(s.todo))):
                values = {}
                if k == 0:
                    values = {1: s.start, 2: str(s.az), 3: s.address, 4: str(s.re), 6: s.anlage}
                elif k == 1:
                    values = {1: "–", 2: "Handy", 3: str(s.zip_code or ""), 4: "Stichtag"}
                elif k == 2:
                    values = {1: s.end, 2: "☐ ja", 3: s.city, 4: s.stichtag}
                values[5] = s.todo[k] if k < len(s.todo) else None
                box_row(values, {1: bold if k != 1 else None, 2: bold if k == 0 else None})
                if k == 1 and s.highlight:
                    ws.cell(z - 1, 5).fill = fill(WARN)
            if s.hint:
                ws.cell(z, 1).value = s.hint
                ws.cell(z, 1).font = arial(size=8, italic=True, color="FF595959")
                z += 1
            if sheet.break_after == i and sheet.break_minutes:
                end = minutes_of(datetime.time.fromisoformat(s.end))
                box_row({1: f"{hh_mm(end)}–{hh_mm(end + sheet.break_minutes)}", 5: "Pause"}, {5: arial(bold=True)})
            if i < len(sheet.stops) - 1 and s.drive_after is not None:
                cell = ws.cell(z, 1)
                cell.value, cell.font, cell.border = f"→ Fahrtzeit: {s.drive_after} min", arial(size=9, italic=True), Border(bottom=THIN)
                z += 1
        cell = ws.cell(z, 1)
        red = sheet.net > MAX_NET_MINUTES or sheet.provisional
        cell.value, cell.font = sheet.net_text(), arial(size=9, bold=True, color="FFC00000" if red else "FF000000")
        z += 2
    _page(ws, landscape=True, last_row=z - 1, margins=dict(left=0.4, right=0.4, top=0.6, bottom=0.6, header=0.3, footer=0.3))
    ws.oddFooter.left.text = name if title else "Fahrpläne"
    ws.oddFooter.right.text = "Seite &P von &N"
    return ws


# --- workbook (excelFahrplaene) -----------------------------------------------------

def sheet_name(text, used):
    """Excel sheet names: max. 31 characters, no \\ / ? * [ ] :, unique."""
    name, k = re.sub(r"[\\/?*\[\]:]", "", text)[:31], 2
    while name in used:
        name = f"{name[:27]} ({k})"
        k += 1
    used.add(name)
    return name


def build_workbook(tours):
    """Workbook + file name for the given tours (same rules as the prototype)."""
    sheets = sorted((tour_sheet(t) for t in tours), key=lambda s: (s.date, s.person))
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    months = sorted({(s.date.year, s.date.month) for s in sheets})
    period = "/".join(dict.fromkeys(MONTHS[m - 1] for _, m in months)) + " " + "/".join(dict.fromkeys(str(y) for y, _ in months))
    if len(sheets) > 1:
        list_sheet(wb, sheet_name("Fahrpläne", used), sheets)
        for person in sorted({s.person for s in sheets}):
            own = [s for s in sheets if s.person == person]
            title = (f"Fahrplan {period} – {person}   ({len({s.date for s in own})} Arbeitstage, {len(own)} Fahrpläne, "
                     f"{own[0].date:%d.%m.} – {own[-1].date:%d.%m.})")
            list_sheet(wb, sheet_name(person, used), own, title)
    for s in sheets:
        day_sheet(wb, s, sheet_name(f"{s.date:%d.%m.} {s.person}", used))
    if len(sheets) == 1:
        filename = f"Fahrplan_{sheets[0].date.isoformat()}_{sheets[0].person}.xlsx"
    else:
        short = "_".join(dict.fromkeys(MONTHS_SHORT[m - 1] for _, m in months))
        filename = f"Fahrplaene_{short}_{months[0][0]}_TomTom_je_Ableser.xlsx"
    return wb, filename
