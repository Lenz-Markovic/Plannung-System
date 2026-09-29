"""Excel export: same structure as the prototype's workbook."""

import datetime
import io

import pytest
from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core import roles
from planning.excel import build_workbook, sheet_name
from planning.models import Tour

pytestmark = pytest.mark.django_db


@pytest.fixture
def demo(monkeypatch):
    monkeypatch.setattr(timezone, "localdate", lambda *args: datetime.date(2026, 9, 28))
    call_command("import_prototype", stdout=io.StringIO())


def reading_tour(min_stops=3):
    return next(t for t in Tour.objects.order_by("date")
                if t.stops.filter(kind="reading").count() >= min_stops and not t.stops.filter(kind="installation").exists())


def test_single_tour_is_one_print_sheet(demo):
    tour = reading_tour()
    workbook, filename = build_workbook([tour])
    assert filename == f"Fahrplan_{tour.date.isoformat()}_{tour.employee}.xlsx"
    assert workbook.sheetnames == [f"{tour.date:%d.%m.} {tour.employee}"]
    ws = workbook.active
    assert ws["A1"].value == "Datum:" and ws["G1"].value == str(tour.employee)
    assert [ws.cell(4, c).value for c in range(1, 9)] == ["Uhrzeit ab:", "AZ", "Adresse", "RE", "ToDo", "FUNK", "Zur Anmeldung", "Termin bestätigt"]
    assert ws.page_setup.orientation == "portrait" and ws.page_setup.paperSize == 9
    assert "VORLÄUFIG" in [c.value for c in ws["A"] if c.value and "Netto" in str(c.value)][0]  # imported = provisional


def test_apartment_access_is_yellow_in_the_todo_cell(demo):
    tour = reading_tour()
    ws = build_workbook([tour])[0].active
    todo_cells = [c for c in ws["E"] if c.value and str(c.value).startswith("Ablesung")]
    highlighted = [c for c in todo_cells if c.fill.fill_type and c.fill.fgColor.rgb == "FFFFE699"]
    assert highlighted and all("🔑 Zugang Wohnung" in c.value for c in highlighted)


def test_several_tours_give_overview_person_and_day_sheets(demo):
    tours = list(Tour.objects.filter(date__month=12)[:4])
    workbook, filename = build_workbook(tours)
    persons = sorted({str(t.employee) for t in tours})
    assert workbook.sheetnames[0] == "Fahrpläne"
    assert workbook.sheetnames[1:1 + len(persons)] == persons
    assert len(workbook.sheetnames) == 1 + len(persons) + len(tours)
    assert filename == "Fahrplaene_Dez_2026_TomTom_je_Ableser.xlsx"
    assert workbook["Fahrpläne"]["A1"].value.startswith("Datum: ")


def test_sheet_names_are_valid_and_unique():
    used = set()
    assert sheet_name("a/b:c?", used) == "abc"
    assert sheet_name("abc", used) == "abc (2)"
    assert len(sheet_name("x" * 40, used)) == 31


def test_download_and_automatic_excel_after_saving(client, demo):
    user = User.objects.create_user(username="d", password="x")
    user.groups.add(Group.objects.get(name=roles.DISPATCHER))
    client.force_login(user)
    tour = reading_tour()
    response = client.get(reverse("planning:tour_excel", args=[tour.pk]))
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response["Content-Disposition"]
    load_workbook(io.BytesIO(response.content))  # a real xlsx file
    page = client.get(reverse("planning:calendar"), {"datum": tour.date.isoformat(), "excel": tour.pk}).content.decode()
    assert f'src="{reverse("planning:tour_excel", args=[tour.pk])}"' in page  # hidden download
    all_tours = client.get(reverse("planning:tours_excel"), {"person": tour.employee_id})
    names = load_workbook(io.BytesIO(all_tours.content)).sheetnames
    assert len(names) == 2 + tour.employee.tours.count()  # overview + person + one sheet per tour


def test_reader_gets_only_own_excel(client, demo):
    user = User.objects.create_user(username="abl", password="x")
    user.groups.add(Group.objects.get(name=roles.READER))
    client.force_login(user)
    assert client.get(reverse("planning:tour_excel", args=[Tour.objects.first().pk])).status_code == 403


# --- montage plans look like reading plans ------------------------------------------

def montage_tour():
    return next(t for t in Tour.objects.order_by("date")
                if t.stops.exists() and not t.stops.filter(kind="reading").exists())


def test_montage_sheet_has_the_same_layout_and_names_the_work(demo):
    tour = montage_tour()
    ws = build_workbook([tour])[0].active
    assert [ws.cell(4, c).value for c in range(1, 9)] == ["Uhrzeit ab:", "AZ", "Adresse", "RE", "ToDo", "FUNK", "Zur Anmeldung", "Termin bestätigt"]
    todo = [c.value for c in ws["E"] if c.value and str(c.value).startswith("Montage RE")]
    assert todo, "the ToDo cell starts with the order number"
    net = next(c.value for c in ws["A"] if c.value and "Netto" in str(c.value))
    assert "(Montage " in net and "Ablesung" not in net
    # the value below the "Stichtag" label comes from the order when there is no building
    stichtag = [ws.cell(c.row + 1, 4).value for c in ws["D"] if c.value == "Stichtag"]
    assert stichtag and "—" not in stichtag


def test_reading_sheet_keeps_the_prototype_wording(demo):
    ws = build_workbook([reading_tour()])[0].active
    net = next(c.value for c in ws["A"] if c.value and "Netto" in str(c.value))
    assert "(Ablesung " in net and "Montage" not in net
