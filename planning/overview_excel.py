"""Excel of "📊 Übersicht" for reports (openpyxl) - the same numbers and filters as the page."""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD = PatternFill("solid", fgColor="1F4E8C")
TOTAL = PatternFill("solid", fgColor="E8EEF7")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
PERCENT = '0" %"'
OUTCOME_COLOURS = ["0CA30C", "FAB219", "D03B3B"]   # ✓ fertig, ◐ teilweise, ✗ nicht erledigt (as on the page)


def _head(ws, row, values, widths=None):
    for column, value in enumerate(values, 1):
        cell = ws.cell(row, column, value)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, width in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row, column).column_letter].width = width


def _line(ws, values, bold=False, total=False):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.border = THIN
        if bold or total:
            cell.font = Font(bold=True)
        if total:
            cell.fill = TOTAL


def _title(ws, text, filters):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = filters
    ws["A2"].font = Font(italic=True, color="555555")


def overview_workbook(data, tiles, filters, user=None):
    """data: overview.collect(...), tiles: overview.headline(...), filters: one line describing the choice."""
    wb = Workbook()
    made = f"{filters} · erstellt {timezone.localtime():%d.%m.%Y %H:%M}" + (f" von {user}" if user else "")

    # 1. Übersicht: the key numbers
    ws = wb.active
    ws.title = "Übersicht"
    _title(ws, "Übersicht Ablesung & Montage", made)
    _head(ws, 4, ["Jetzt offen", "Anzahl"], [40, 14])
    for label, value in [("Rückmeldungen offen (im Büro zu bearbeiten)", tiles["followup_open"]),
                         ("Nachtermin nötig (Liegenschaften + Aufträge)", tiles["revisits"]),
                         ("Über der 14-Tage-Frist", tiles["overdue"]), ("bald fällig (14-Tage-Frist)", tiles["soon"]),
                         ("Konflikte offen", tiles["conflicts"]), ("Stopps heute", tiles["planned_today"])]:
        _line(ws, [label, value])
    ws.append([])
    start = ws.max_row + 1
    _head(ws, start, ["Termin-Ergebnisse im Zeitraum", "Anzahl", "Anteil"])
    ws.column_dimensions["C"].width = 12
    for key, label, n, share in data["totals"]:
        _line(ws, [label, n, share])
        ws.cell(ws.max_row, 3).number_format = PERCENT
    _line(ws, ["gemeldet zusammen", data["reported"], ""], total=True)
    ws.append([])
    start = ws.max_row + 1
    _head(ws, start, ["Liegenschaften nach Status", "Anzahl", "Anteil"])
    for key, label, n, share in data["buildings"]:
        _line(ws, [label, n, round(share)])
        ws.cell(ws.max_row, 3).number_format = PERCENT
    _line(ws, ["Liegenschaften zusammen", data["buildings_total"], ""], total=True)
    ws.append([])
    start = ws.max_row + 1
    _head(ws, start, ["Montageaufträge nach Status", "Anzahl"])
    for key, label, n in data["orders"]:
        _line(ws, [label, n])
    _line(ws, ["Aufträge zusammen", data["orders_total"]], total=True)

    # 2. Termine pro Tag / Woche (+ a stacked column chart)
    ws = wb.create_sheet("Termine je Woche" if data["per_week"] else "Termine je Tag")
    _title(ws, "Termin-Ergebnisse " + ("pro Woche" if data["per_week"] else "pro Tag"), made)
    _head(ws, 4, ["Woche" if data["per_week"] else "Tag", "✓ fertig", "◐ teilweise", "✗ nicht erledigt", "zusammen"],
          [12, 12, 13, 16, 13])
    first = ws.max_row + 1
    for c in data["columns"]:
        b = c["bucket"]
        _line(ws, [b.label, b.complete, b.partial, b.absent, b.total])
    last = ws.max_row
    _line(ws, ["Gesamt", *[n for _, _, n, _ in data["totals"]], data["reported"]], total=True)
    ws.freeze_panes = "A5"
    if last >= first:
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title = "Termin-Ergebnisse"
        chart.y_axis.title = "Termine"
        chart.height, chart.width = 9, 22
        chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=4, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        for series, colour in zip(chart.series, OUTCOME_COLOURS):
            series.graphicalProperties.solidFill = colour
            series.graphicalProperties.line.solidFill = colour
        ws.add_chart(chart, "G4")

    # 3. Nach Stichtag (Abrechnungszeitraum)
    ws = wb.create_sheet("Nach Stichtag")
    _title(ws, "Liegenschaften nach Stichtag (Abrechnungszeitraum)", made)
    _head(ws, 4, ["Stichtag", "Abrechnung von", "bis", "Liegenschaften", "offen", "Nacharbeit", "freigegeben",
                  "% freigegeben", "geplant", "ohne Termin", "Nachtermin nötig"], [12, 16, 12, 16, 9, 13, 14, 15, 10, 13, 18])
    for r in data["stichtag_rows"]:
        _line(ws, [r.stichtag, r.period_start, r.period_end, r.total, r.open, r.rework, r.released, r.released_share,
                   r.planned, r.unplanned, r.revisit])
        line = ws.max_row
        for column in (1, 2, 3):
            ws.cell(line, column).number_format = "DD.MM.YYYY"
        ws.cell(line, 8).number_format = PERCENT
    rows = data["stichtag_rows"]
    if len(rows) > 1:
        _line(ws, ["Gesamt", "", "", *[sum(getattr(r, f) for r in rows) for f in ("total", "open", "rework", "released")],
                   "", *[sum(getattr(r, f) for r in rows) for f in ("planned", "unplanned", "revisit")]], total=True)
    ws.freeze_panes = "A5"

    # 4. Pro Person
    ws = wb.create_sheet("Pro Person")
    _title(ws, "Pro Person / Team – Stopps in den Fahrplänen des Zeitraums (bis heute)", made)
    _head(ws, 4, ["Person / Team", "Stopps geplant", "davon gemeldet", "% gemeldet", "✓ fertig", "◐ teilweise",
                  "✗ nicht erledigt", "Erfolg", "noch offen"], [24, 13, 13, 11, 10, 11, 15, 10, 11])
    for p in data["people"]:
        _line(ws, [p.people, p.planned, p.reported, p.reported_share, p.complete, p.partial, p.absent,
                   p.success if (p.complete or p.partial or p.absent) else None, p.open])
        ws.cell(ws.max_row, 4).number_format = ws.cell(ws.max_row, 8).number_format = PERCENT
    people = data["people"]
    if people:
        _line(ws, ["Gesamt", *[sum(getattr(p, f) for p in people) for f in ("planned", "reported")], "",
                   *[sum(getattr(p, f) for p in people) for f in ("complete", "partial", "absent")], "",
                   sum(p.open for p in people)], total=True)
    ws.freeze_panes = "A5"

    # 5. Gründe und wie viele Termine bis fertig
    ws = wb.create_sheet("Gründe & Termine")
    _title(ws, "Warum nicht erledigt? · Wie viele Termine bis fertig?", made)
    _head(ws, 4, ["Grund bei ✗ nicht erledigt", "Anzahl"], [40, 10])
    for label, n, _ in data["reasons"]:
        _line(ws, [label, n])
    if not data["reasons"]:
        _line(ws, ["kein ✗ im Zeitraum", 0])
    ws.append([])
    start = ws.max_row + 1
    _head(ws, start, ["✓ fertig beim …", "Anzahl"])
    for label, n, _ in data["attempts"]:
        _line(ws, [f"{label} Termin", n])
    _line(ws, ["✓ fertig zusammen", data["attempts_total"]], total=True)

    for sheet in wb.worksheets:  # printing: landscape, one page wide
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_options.gridLines = False

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
