"""📄 Aushänge per Fahrplan (what is still to print) and the 🗺 Aushang-Route (a printed list)."""

import datetime
from io import BytesIO

import pytest
from django.contrib.auth.models import Group, User
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from core import roles
from documents import notice_overview, notices
from documents import notice_rules as rules
from planning import services
from planning.models import StopKind, Tour, TourStop

pytestmark = pytest.mark.django_db
TODAY = datetime.date(2026, 9, 28)


@pytest.fixture
def demo(demo_import, monkeypatch):
    monkeypatch.setattr(timezone, "localdate", lambda *args: TODAY)
    monkeypatch.setattr(services, "get_client", lambda: None)


def login(role, name):
    user = User.objects.create_user(username=name)
    user.groups.add(Group.objects.get(name=role))
    client = Client()
    client.force_login(user)
    return client


def a_tour():
    """A coming plan with at least 2 readings (imported: no saved times -> the time is estimated)."""
    for tour in Tour.objects.filter(date__gt=TODAY).order_by("date", "pk"):
        if tour.stops.filter(kind=StopKind.READING).count() >= 2:
            return tour, list(tour.stops.filter(kind=StopKind.READING).order_by("position"))
    raise AssertionError("no plan")


def test_page_shows_every_plan_and_what_is_still_to_print(demo):
    c = login(roles.DISPATCHER, "dispo")
    tour, stops = a_tour()
    html = c.get(reverse("documents:aushaenge"), {"f": "alle", "zeitraum": "alle"}).content.decode()
    assert "📄 Aushänge" in html and tour.people_label in html and stops[0].building.file_number in html
    assert "kein Aushang" in html and "geschätzt" in html and "aushängen bis" in html

    answer = c.post(reverse("documents:aushang_wanted", args=[stops[0].pk]), {"on": "1"}).content.decode()
    assert f'id="ah-head-{tour.pk}" hx-swap-oob="true"' in answer and "1 fehlen" in answer   # the head counts again
    block = next(b for b in notice_overview.blocks(TODAY, None) if b.tour.pk == tour.pk)
    assert [s.stop.pk for s in block.missing] == [stops[0].pk]
    only_open = c.get(reverse("documents:aushaenge"), {"zeitraum": "alle"}).content.decode()
    assert f'id="ah-{stops[0].pk}"' in only_open and "noch nicht gedruckt" in only_open

    # print it (the print page marks it) -> printed, the plan is no longer "noch zu drucken"
    c.post(reverse("documents:notice_print"), {"stop": [stops[0].pk]})
    block = next(b for b in notice_overview.blocks(TODAY, None) if b.tour.pk == tour.pk)
    assert block.printed and not block.missing
    assert tour.pk not in [b.tour.pk for b in notice_overview.blocks(TODAY, None, only_open=True)]


def test_briefe_for_single_flats_one_page_each(demo):
    c = login(roles.PROCESSING, "sb")
    tour, stops = a_tour()
    stop = stops[0]
    html = c.post(reverse("documents:aushang_edit", args=[stop.pk]),
                  {"scope": rules.SOME_UNITS, "units": "Whg 3 (Müller), Whg 7", "von": "8", "bis": "12"}).content.decode()
    assert "2 Briefe" in html and "Whg 3 (Müller), Whg 7" in html and "von Hand" in html and "08:00–12:00 Uhr" in html
    stop.refresh_from_db()
    assert stop.notice_wanted and stop.notice_from == datetime.time(8)
    pages = notices.pages_of(stop)
    assert [flat for flat, _ in pages] == ["Whg 3 (Müller)", "Whg 7"]
    assert pages[0][1].bottom == "Für: Whg 3 (Müller)" and pages[0][1].time == "08:00 – 12:00 Uhr"
    page = c.get(reverse("documents:notice_page"), {"stop": stop.pk}).content.decode()
    assert page.count('class="page"') == 2 and "Brief für Whg 7" in page
    bad = c.post(reverse("documents:aushang_edit", args=[stop.pk]), {"von": "25:99"})
    assert bad["HX-Reswap"] == "none"
    services.save_draft(services.draft_from_tour(Tour.objects.get(pk=tour.pk)), None, confirm=False)  # saved again
    again = TourStop.objects.get(tour_id=tour.pk, building_id=stop.building_id, kind=StopKind.READING)
    assert (again.notice_scope, again.notice_units, again.notice_from) == (rules.SOME_UNITS, "Whg 3 (Müller), Whg 7",
                                                                          datetime.time(8))


def test_route_from_and_to_the_office_without_a_fixed_time(demo):
    from openpyxl import load_workbook

    c = login(roles.DISPATCHER, "dispo")
    tour, stops = a_tour()
    for s in stops:
        c.post(reverse("documents:aushang_wanted", args=[s.pk]), {"on": "1"})
    ids = [s.pk for s in stops]
    html = c.get(reverse("documents:aushang_route"), {"stop": ids}).content.decode()
    assert "🗺 Aushang-Route" in html and "🖨 Route drucken" in html and stops[0].building.street in html
    assert "Zuckerfabrik 14, 70376 Stuttgart" in html and "zurück ins Büro" in html and "nach" in html
    found, back, office = notice_overview.route(ids)
    houses = {s.building_id for s in stops}
    assert len(found) == len(houses) and [s.n for s in found] == list(range(1, len(found) + 1))
    assert found[0].arrive == found[0].drive_minutes                     # counted from leaving the office (0:00)
    assert all(a.leave <= b.arrive for a, b in zip(found, found[1:])) and back.arrive >= found[-1].leave
    assert sum(len(s.papers) for s in found) == len(stops) and found[0].minutes == 4
    timed, _, _ = notice_overview.route(ids, datetime.date(2026, 10, 2), datetime.time(9))
    assert timed[0].arrive == 9 * 60 + timed[0].drive_minutes
    assert "an" in c.get(reverse("documents:aushang_route"), {"stop": ids, "ab": "9:00"}).content.decode()
    xlsx = c.get(reverse("documents:aushang_route"), {"stop": ids, "format": "xlsx"})
    sheet = load_workbook(BytesIO(xlsx.content)).active
    assert sheet["A1"].value == "Aushang-Route" and sheet.cell(5, 1).value == "B" and sheet.cell(6, 1).value == 1
    rows = list(sheet.iter_rows(values_only=True))
    assert any(str(r[1]).startswith("zurück ins Büro") for r in rows) and any(r[1] == "Gesamt" for r in rows)
    assert "Keine Häuser gewählt" in c.get(reverse("documents:aushang_route")).content.decode()


def test_printed_is_ready_to_hand_out(demo):
    c = login(roles.DISPATCHER, "dispo2")
    tour, stops = a_tour()
    c.post(reverse("documents:aushang_wanted", args=[stops[0].pk]), {"on": "1"})
    c.post(reverse("documents:notice_print"), {"stop": [stops[0].pk]})
    html = c.get(reverse("documents:aushaenge"), {"f": "bereit", "zeitraum": "alle"}).content.decode()
    assert "gedruckt – bereit zum Verteilen" in html and f'id="ah-{stops[0].pk}"' in html


def test_roles(demo):
    tour, stops = a_tour()
    lead = login(roles.MANAGEMENT, "leitung")
    assert lead.get(reverse("documents:aushaenge")).status_code == 200
    assert lead.post(reverse("documents:aushang_wanted", args=[stops[0].pk]), {"on": "1"}).status_code == 403
    assert login(roles.READER, "abl").get(reverse("documents:aushaenge")).status_code == 403
    assert "📄 Aushänge &amp; Aushang-Route" in login(roles.PROCESSING, "sb2").get(reverse("planning:calendar")).content.decode()
