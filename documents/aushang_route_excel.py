"""Excel of the 🗺 Aushang-Route: one row per house, distance / drive to the next one, totals (openpyxl)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD = PatternFill("solid", fgColor="1F4E8C")
START = PatternFill("solid", fgColor="D9EAD3")     # green
STOP = PatternFill("solid", fgColor="DDE7F6")      # blue
END = PatternFill("solid", fgColor="F4CCCC")       # red
WARN = PatternFill("solid", fgColor="FCE5B3")      # amber: warning
TOTAL = PatternFill("solid", fgColor="E8EEF7")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))


def _clock(minutes):
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def route_workbook(context):
    route, back, office = context["route"], context["back"], context["office"]
    has_clock = context["has_clock"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Aushang-Route"
    title = "Aushang-Route" + (f" {context['day']:%d.%m.%Y}" if context.get("day") else "")
    ws["A1"] = title + (f" · {context['person']}" if context.get("person") else "")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Start und Ende: Büro, {office[0]} · {len(route)} Häuser · {context['papers']} Aushänge/Briefe · "
                f"Dauer ca. {context['total']} · Fahrzeiten {'TomTom' if context.get('all_tomtom') else 'geschätzt'} · "
                f"Reihenfolge {'TomTom (echte Straßen)' if back.order_source == 'tomtom' else 'nach Luftlinie'}")
    ws["A2"].font = Font(italic=True, color="555555")
    head = ["Nr.", "Straße", "PLZ / Ort", "Ankunft" if has_clock else "nach", "Was", "für Termin", "km zum nächsten",
            "Fahrt zum nächsten (min)", "Hinweise", "✓"]
    ws.append([])
    ws.append(head)
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def when(minutes):
        return _clock(minutes) if has_clock else ("Start" if minutes == 0 else f"+{minutes} min")

    def row(values, fill, notes=None):
        ws.append(values)
        line = ws.max_row
        for cell in ws[line]:
            cell.border, cell.fill = THIN, fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if notes:
            ws.cell(line, 9).fill = WARN

    first = route[0] if route else None
    start = (context["start"].hour * 60 + context["start"].minute) if has_clock else 0
    street, _, place = office[0].partition(", ")
    row(["B", f"Büro – {street}", place, when(start), "Start", "",
         float(first.drive_km) if first and first.drive_km is not None else None, first.drive_minutes if first else None,
         "", ""], START)
    for i, s in enumerate(route):
        following = route[i + 1] if i + 1 < len(route) else None
        nxt_km, nxt_min = ((following.drive_km, following.drive_minutes) if following else (back.km, back.minutes))
        notes = []
        if s.point_source != "tomtom":
            notes.append("Adresse nur ungefähr (PLZ-Mitte)" if s.point else "Adresse ohne Position")
        if not (following.drive_from_tomtom if following else back.from_tomtom):
            notes.append("Fahrzeit geschätzt")
        if s.far_km:
            notes.append(f"liegt weit weg (nächstes Haus {s.far_km} km)")
        row([s.n, s.target.street, f"{s.target.zip_code} {s.target.city}", when(s.arrive),
             "; ".join(f"{p[0]}: {p[1]}" for p in s.papers),
             ", ".join(sorted({f"{p[2]:%d.%m.} {p[3]}" for p in s.papers})),
             float(nxt_km) if nxt_km is not None else None, nxt_min, ", ".join(notes), "☐"], STOP, notes)
    row(["B", f"zurück ins Büro – {street}", place, when(back.arrive), "Ende", "", None, None, "", ""], END)
    ws.append(["", "Gesamt", f"{len(route)} Häuser", context["total"], f"{context['papers']} Aushänge/Briefe", "",
               float(context["total_km"]), context["total_drive"], f"Verteilen ca. {context['total_work']} min", ""])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = Font(bold=True), TOTAL
    for column, width in zip("ABCDEFGHIJ", [5, 30, 22, 9, 42, 20, 10, 12, 28, 4]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
